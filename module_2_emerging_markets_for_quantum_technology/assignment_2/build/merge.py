import os, sys
HERE = os.path.dirname(os.path.abspath(__file__))
ASSIGN_DIR = os.path.dirname(HERE)
OUT_DIR = os.environ.get('OUTDIR', ASSIGN_DIR)
sys.path.insert(0, HERE)
import copy
from docx import Document
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK
from references import REFERENCES, parse_italic_segments

MON = f"{OUT_DIR}/MON_filled.docx"
VPC = f"{OUT_DIR}/VPC_filled.docx"
BMC = f"{OUT_DIR}/BMC_filled.docx"
OUT = f"{OUT_DIR}/Assessment2_Quantum_Portfolio.docx"

try:
    from docxcompose.composer import Composer
    HAVE_COMPOSER = True
except ImportError:
    HAVE_COMPOSER = False


# Student name, placed at the left of the template's running head. Change it here,
# not by hand in Word -- an edit made in the .docx is lost on the next rebuild.
NAME = "Hung Han"
RUNNING_HEAD = "Quantum Technology"   # right-hand side, supplied by the Sussex template

TITLE = "Assessment 2 Portfolio"
SUBTITLE = ("A proposed UK venture commercialising cold-atom gravity gradiometry "
            "for subsurface survey.")


def add_title_block(doc):
    """Insert title + subtitle before the first paragraph of the merged document."""
    first = doc.paragraphs[0]
    sub = first.insert_paragraph_before(SUBTITLE)
    try:
        sub.style = doc.styles['Normal']
    except KeyError:
        pass
    ttl = sub.insert_paragraph_before(TITLE)
    try:
        ttl.style = doc.styles['Title']
    except KeyError:
        pass


def add_name_to_header(doc):
    """Put NAME at the left of the running head, keeping the template's module title
    right-aligned on the same line via a right tab stop at the right margin."""
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
    done = 0
    for section in doc.sections:
        for p in section.header.paragraphs:
            if p.text.strip() != RUNNING_HEAD or not p.runs:
                continue
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            width = section.page_width - section.left_margin - section.right_margin
            p.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
            first_r = p.runs[0]._r
            new_r = copy.deepcopy(first_r)          # inherit the SSX_RunningHead run formatting
            for t in list(new_r.findall(qn('w:t'))):
                new_r.remove(t)
            t = new_r.makeelement(qn('w:t'), {})
            t.text = NAME
            new_r.append(t)
            new_r.append(new_r.makeelement(qn('w:tab'), {}))
            first_r.addprevious(new_r)
            done += 1
    if not done:
        raise RuntimeError(f"running head {RUNNING_HEAD!r} not found; name not inserted")


def add_page_break(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)


def add_references_section(doc):
    add_page_break(doc)
    doc.add_heading('References', level=1)
    for ref in REFERENCES:
        p = doc.add_paragraph()
        for text, italic in parse_italic_segments(ref):
            run = p.add_run(text)
            run.italic = italic


if HAVE_COMPOSER:
    master = Document(MON)
    add_page_break(master)
    composer = Composer(master)
    composer.append(Document(VPC))
    master_doc_after_vpc = master  # composer mutates `master` in place
    add_page_break(master)
    composer.append(Document(BMC))
    add_references_section(master)
    add_title_block(master)
    add_name_to_header(master)
    composer.save(OUT)
    print(f"Merged via docxcompose -> {OUT}")
else:
    # Fallback: manual body-element concatenation
    master = Document(MON)
    add_page_break(master)

    def append_body(master_doc, src_path):
        src = Document(src_path)
        src_body = src.element.body
        for el in list(src_body):
            if el.tag == qn('w:sectPr'):
                continue
            master_doc.element.body.append(copy.deepcopy(el))

    append_body(master, VPC)
    add_page_break(master)
    append_body(master, BMC)
    add_references_section(master)
    add_title_block(master)
    add_name_to_header(master)
    master.save(OUT)
    print(f"Merged via fallback body-concat -> {OUT}")
