import os, sys
HERE = os.path.dirname(os.path.abspath(__file__))
ASSIGN_DIR = os.path.dirname(HERE)
OUT_DIR = os.environ.get('OUTDIR', ASSIGN_DIR)
sys.path.insert(0, HERE)
from docx import Document
from docx.oxml.ns import qn
from helpers import (fix_typography, set_checkbox, find_checkbox_in_paragraph,
                      paragraph_label_text, clone_block, insert_block_after,
                      add_run_after_label, new_paragraph)

SRC = f"{ASSIGN_DIR}/Market Opportunity Navigator_Blank.docx"
OUT = f"{OUT_DIR}/MON_filled.docx"

T = lambda s: fix_typography(s)

# ---------- content ----------

CA = {
    1: T('Measures the gravitational difference between two clouds of laser-cooled atoms falling in one apparatus. It reads density contrast, so it needs no line of sight, no contact and no excavation.'),
    2: T('Vibration is common to both clouds and cancels differentially; accelerometers and shielding handle the rest. This is what took the instrument off the optical bench and onto a street (Stray et al., 2022).'),
    3: T("Bayesian inversion turns sparse gradient readings into a three-dimensional "
         "density estimate, and reports its confidence in position and depth rather than "
         "drawing one certain line."),
    4: T('Packaging lasers, vacuum and shielding into equipment a trained surveyor can run through a working day. Most of it is tacit knowledge held by people, which matters for value capture.'),
}

OPPS_PART1 = {
    1: dict(
        application=T("Targeted pre-excavation survey of high-risk urban sites, finding "
                       "buried assets, voids and abandoned structures that electromagnetic "
                       "location and radar miss."),
        customer_group=T('UK water, gas, power and telecoms asset owners, the tier-1 contractors digging for them, and highway and rail authorities.'),
    ),
    2: dict(
        application=T('Density targeting for mineral exploration and geothermal assessment, narrowing where to drill before anyone drills.'),
        customer_group=T("Exploration companies, geothermal developers, and the "
                          "geophysical survey contractors who serve them."),
    ),
    3: dict(
        application=T("Drift-free inertial navigation where satellite signals are "
                       "jammed, spoofed or unavailable."),
        customer_group=T("The Ministry of Defence and defence primes first, maritime "
                          "and aviation operators later."),
    ),
}

OPPS_PART2 = {
    1: dict(
        name="Opportunity 1: Pre-excavation subsurface survey",
        criteria={
            'Problem Severity': 'HIGH', 'Pertinent Solution': 'HIGH', 'Impact Reach': 'HIGH',
            'Overall Impact': 'HIGH',
            'Compelling Reasons to Buy': 'HIGH', 'Potential Market Volume': 'MEDIUM',
            'Economic Viability': 'HIGH', 'Overall Potential': 'HIGH',
            'Implementation Obstacles': 'MEDIUM', 'Time to Revenue': 'MEDIUM',
            'External Risks': 'HIGH', 'Overall Challenge': 'MEDIUM',
        },
        # one Notes: paragraph under each Overall score, plus one under the category
        notes={
            'Overall Impact': T(
                'Notes: One hole in every 65 hits something: about 60,000 strikes a '
                'year and GBP 2.4bn in economic cost (Utility Strike Avoidance Group, '
                '2023). Workers are hurt.'),
            'Overall Potential': T(
                'Notes: Contractors already buy PAS 128 surveys, so a budget line and a '
                'standard exist to sell against (British Standards Institution, 2022), '
                'and reported strike rates mean improvement can be shown rather than '
                'claimed.'),
            'Overall Challenge': T(
                'Notes: The physics already works outdoors, so this is productisation, '
                'not discovery. Two risks bite: PAS 128 recognises radar and '
                'electromagnetic location, not gravity; and Delta.g, a Birmingham '
                'spin-out with GBP 4.6m of seed funding, is ahead of us (The Quantum '
                'Insider, 2025).'),
            'Opportunity Category': T(
                'Notes: Gold Mine: highest potential, lowest challenge of the three, '
                'judged against each other, not in the abstract.'),
        },
        category='GOLD MINE',
    ),
    2: dict(
        name="Opportunity 2: Mineral and geothermal exploration",
        criteria={
            'Problem Severity': 'MEDIUM', 'Pertinent Solution': 'MEDIUM', 'Impact Reach': 'MEDIUM',
            'Overall Impact': 'MEDIUM',
            'Compelling Reasons to Buy': 'HIGH', 'Potential Market Volume': 'HIGH',
            'Economic Viability': 'MEDIUM', 'Overall Potential': 'HIGH',
            'Implementation Obstacles': 'HIGH', 'Time to Revenue': 'HIGH',
            'External Risks': 'MEDIUM', 'Overall Challenge': 'HIGH',
        },
        notes={
            'Overall Impact': T(
                'Notes: Better targeting means fewer speculative boreholes, though the '
                'harm avoided is diffuse next to a struck gas main.'),
            'Overall Potential': T(
                'Notes: Global exploration budgets dwarf UK survey spend, and drilling '
                'costs enough that better targets are worth paying for.'),
            'Overall Challenge': T(
                'Notes: Remote sites, power and calibration in harsh conditions sit '
                'beyond our packaging. Campaigns are seasonal, procurement slow, and '
                'airborne gradiometry already serves it.'),
        },
        category='MOON SHOT',
    ),
    3: dict(
        name="Opportunity 3: GPS-denied navigation",
        criteria={
            'Problem Severity': 'HIGH', 'Pertinent Solution': 'MEDIUM', 'Impact Reach': 'MEDIUM',
            'Overall Impact': 'MEDIUM',
            'Compelling Reasons to Buy': 'HIGH', 'Potential Market Volume': 'MEDIUM',
            'Economic Viability': 'MEDIUM', 'Overall Potential': 'HIGH',
            'Implementation Obstacles': 'VERY HIGH', 'Time to Revenue': 'VERY HIGH',
            'External Risks': 'HIGH', 'Overall Challenge': 'VERY HIGH',
        },
        notes={
            'Overall Impact': T(
                'Notes: Satellite navigation is jammed and spoofed routinely, with '
                'consequences from delayed shipping to lost life. The national mission '
                'targets quantum navigation on aircraft by 2030 (HM Government, 2023), '
                'so the state agrees the problem is severe even without a commercial '
                'market.'),
            'Overall Potential': T(
                'Notes: Defence buyers pay well, but are few and slow.'),
            'Overall Challenge': T(
                'Notes: A survey instrument may stand still; a navigator must work '
                'while moving, at a fraction of the size, weight and power. Export '
                'control applies, and the dual-use question becomes ours to answer.'),
        },
        category='MOON SHOT',
    ),
}

PRIMARY_OPP = T("Opportunity 1, pre-excavation subsurface survey.")
OTHER_OPPS = T("Opportunity 2, mineral and geothermal exploration; Opportunity 3, "
               "GPS-denied navigation.")
OPP1_AGILE_FALLBACK = "Opportunity 1 — Agile Strategy: Pursue Now"

OPP_ASSESS = {
    2: dict(
        heading="Opportunity 2 — Mineral and geothermal exploration",
        product_relatedness='HIGH', market_relatedness='LOW',
        backup='YES', growth='NO', agile='KEEP OPEN',
    ),
    3: dict(
        heading="Opportunity 3 — GPS-denied navigation",
        product_relatedness='MEDIUM', market_relatedness='LOW',
        backup='NO', growth='YES', agile='KEEP OPEN',
    ),
}

NOTES_TEXT = T(
    ("Notes: The kernel is Rumelt's (2011): our diagnosis is a laboratory-"
        'grade sensor with no route to a customer, so the guiding policy is to '
        'earn revenue surveying the ground everyone else struggles with, maturing'
        " the instrument on a client's site rather than in our own. Exploration "
        'is a genuine back-up, its buyers unrelated to UK construction, so a '
        'stall in contractor adoption would not stop it. It is not a hedge '
        'against technical failure: all three need the same sensor to work, and '
        'pretending otherwise is the fiction this tool exists to prevent (Gruber '
        'and Tal, 2017). Keeping both open costs design headroom and a data-'
        'sharing agreement, but no engineers.'))

# ---------- helper: generic checkbox-group scanner ----------

def get_style_id(p_el):
    pPr = p_el.find(qn('w:pPr'))
    if pPr is None:
        return None
    pStyle = pPr.find(qn('w:pStyle'))
    if pStyle is None:
        return None
    return pStyle.get(qn('w:val'))


def scan_groups(block_els):
    """Return dict: normalized_label -> (heading_el, [(option_text, sdt_el), ...])"""
    groups = {}
    n = len(block_els)
    i = 0
    while i < n:
        el = block_els[i]
        style = get_style_id(el)
        if style in ('Heading1', 'Heading2', 'Heading3'):
            label = paragraph_label_text(el).strip().rstrip(':').strip()
            j = i + 1
            options = []
            while j < n:
                sub = block_els[j]
                substyle = get_style_id(sub)
                if substyle in ('Heading1', 'Heading2', 'Heading3'):
                    break
                cb = find_checkbox_in_paragraph(sub)
                if cb is not None:
                    opt_text = paragraph_label_text(sub).strip()
                    options.append((opt_text, cb))
                j += 1
            if options:
                groups[label] = (el, options)
            i = j
        else:
            i += 1
    return groups


def tick(options, value):
    value_norm = value.strip().upper()
    found = False
    for opt_text, sdt in options:
        opt_norm = opt_text.strip().upper()
        is_match = opt_norm == value_norm or (
            opt_norm.startswith(value_norm)
            and (len(opt_norm) == len(value_norm) or opt_norm[len(value_norm)] in ' (')
        )
        set_checkbox(sdt, is_match)
        if is_match:
            found = True
    if not found:
        raise ValueError(f"Value {value!r} not found among options "
                          f"{[o for o,_ in options]}")


# ---------- main ----------

doc = Document(SRC)
paras = doc.paragraphs

# -- Part 1 --
add_run_after_label(paras[6]._p, CA[1])
add_run_after_label(paras[8]._p, CA[2])
add_run_after_label(paras[10]._p, CA[3])
add_run_after_label(paras[12]._p, CA[4])

add_run_after_label(paras[16]._p, OPPS_PART1[1]['application'])
add_run_after_label(paras[17]._p, OPPS_PART1[1]['customer_group'])
add_run_after_label(paras[19]._p, OPPS_PART1[2]['application'])
add_run_after_label(paras[20]._p, OPPS_PART1[2]['customer_group'])
add_run_after_label(paras[22]._p, OPPS_PART1[3]['application'])
add_run_after_label(paras[23]._p, OPPS_PART1[3]['customer_group'])

# -- Part 2: duplicate the block (idx 27..121 inclusive) so there are 3 copies --
block_start, block_end = 27, 121
orig_block_els = [paras[i]._p for i in range(block_start, block_end + 1)]

clone2 = clone_block(orig_block_els)
last = insert_block_after(orig_block_els[-1], clone2)
clone3 = clone_block(orig_block_els)
insert_block_after(last, clone3)

blocks = [orig_block_els, clone2, clone3]

for opp_num, block_els in zip([1, 2, 3], blocks):
    data = OPPS_PART2[opp_num]
    # Opportunity Name: is the first paragraph in the block
    name_p = block_els[0]
    assert paragraph_label_text(name_p).strip() == 'Opportunity Name:', \
        paragraph_label_text(name_p)
    add_run_after_label(name_p, data['name'])

    groups = scan_groups(block_els)
    for crit, val in data['criteria'].items():
        heading_el, options = groups[crit]
        tick(options, val)

    cat_heading_el, cat_options = groups['Opportunity Category']
    tick(cat_options, data['category'])

    # insert a Notes: paragraph directly after the last checkbox of each group it
    # explains (Overall Impact / Potential / Challenge, and the category), so the
    # rationale sits beside the score rather than at the foot of the card
    for group_label, text in data['notes'].items():
        _, options = groups[group_label]
        last_option_el = options[-1][1].getparent()  # the <w:p> containing the sdt
        last_option_el.addnext(new_paragraph('Normal', text))

# -- Part 3 --
# refresh paragraph list / locate by scanning doc.paragraphs again since indices shifted
doc_paras = doc.paragraphs
primary_p = None
other_p = None
assess_heading_p = None
notes_anchor_p = None
for p in doc_paras:
    txt = p.text.strip()
    if txt == 'Primary Opportunity:':
        primary_p = p
    elif txt == 'Other Opportunities:':
        other_p = p
    elif txt == 'Opportunity Assessment' and assess_heading_p is None:
        assess_heading_p = p

add_run_after_label(primary_p._p, PRIMARY_OPP)
add_run_after_label(other_p._p, OTHER_OPPS)

# Opportunity 1 fallback short paragraph, right after "Other Opportunities:" paragraph
p_opp1 = new_paragraph('Normal', OPP1_AGILE_FALLBACK)
other_p._p.addnext(p_opp1)

# Locate the single "Opportunity Assessment" block: heading (Heading2) through the
# last checkbox group ("Agile Strategy" options), i.e. up to end of doc body content
# paragraphs (before the two trailing empty paragraphs).
all_p_els = list(doc.element.body.findall(qn('w:p')))
start_idx = all_p_els.index(assess_heading_p._p)
# find end: scan groups from start_idx to find Agile Strategy group's last option paragraph
tail_els = all_p_els[start_idx:]
groups3 = scan_groups(tail_els)
agile_heading_el, agile_options = groups3['Agile Strategy']
end_el = agile_options[-1][1].getparent()
end_idx = all_p_els.index(end_el)
assess_block_els = all_p_els[start_idx:end_idx + 1]

# the template already has ONE instance of this block; duplicate it ONCE more so the
# block appears twice in total: original -> Opportunity 2, clone -> Opportunity 3.
clone_o3 = clone_block(assess_block_els)
insert_block_after(assess_block_els[-1], clone_o3)

assess_blocks = {2: assess_block_els, 3: clone_o3}

for opp_num, block_els in assess_blocks.items():
    data = OPP_ASSESS[opp_num]
    heading_el = block_els[0]
    label_p = new_paragraph('Heading3', data['heading'])
    heading_el.addprevious(label_p)

    groups = scan_groups(block_els)
    _, pr_opts = groups['Product Relatedness']
    tick(pr_opts, data['product_relatedness'])
    _, mr_opts = groups['Market Relatedness']
    tick(mr_opts, data['market_relatedness'])
    _, bu_opts = groups['Back-Up Option?']
    tick(bu_opts, data['backup'])
    _, go_opts = groups['Growth Option?']
    tick(go_opts, data['growth'])
    _, ag_opts = groups['Agile Strategy']
    tick(ag_opts, data['agile'])

# Add a "Notes:" paragraph at the very end (template has none) with the Part 3 notes text.
last_body_p = list(doc.element.body.findall(qn('w:p')))[-1]
# the last two paragraphs in the template are empty trailing paragraphs; insert Notes
# before them is fine, but simplest: append at very end before sectPr, i.e. after the
# last <w:p>.
notes_p = new_paragraph('Normal', NOTES_TEXT)
last_body_p.addnext(notes_p)

doc.save(OUT)
print("Saved", OUT)
